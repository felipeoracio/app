"""Validated text-layer extraction for UnoWord document uploads.

Phase 16 — pluggable document pipeline. Each supported format is a
`DocumentProcessor` registered in a small registry; the public
`extract_document()` only dispatches by file extension. Adding a new format
later (e.g. RTF, or an OCR/image processor) is a matter of subclassing
`DocumentProcessor` and calling `register_processor(...)` — no change to the
dispatcher, the router, or the storage/chunking/embedding flow. Example:

    class ImageOcrProcessor(DocumentProcessor):
        name = "image-ocr"
        mime_by_extension = {".png": "image/png", ".jpg": "image/jpeg"}
        def extract(self, data: bytes) -> str:
            return run_ocr(data)   # supplied by a future OCR service
    register_processor(ImageOcrProcessor())

OCR / scanned-image support is intentionally NOT implemented here (no such
infrastructure exists yet); images currently fall through to the unsupported
path with a message pointing at a future OCR processor.
"""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_DOCX_FILES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024

# Populated by register_processor() at import time. Kept as a module-level dict
# so existing importers (routers, tests) keep working unchanged.
MIME_BY_EXTENSION: dict[str, str] = {}


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    file_type: str
    text: str
    word_count: int


class DocumentProcessingError(ValueError):
    pass


def count_words(text: str) -> int:
    return sum(1 for token in (text or "").split() if any(character.isalnum() for character in token))


def _extract_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentProcessingError("TXT and Markdown files must use UTF-8 encoding") from exc


def _extract_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise DocumentProcessingError("The file does not contain a valid PDF header")
    try:
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise DocumentProcessingError("Password-protected PDFs are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentProcessingError(f"PDFs are limited to {MAX_PDF_PAGES} pages")
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("The PDF could not be read safely") from exc


def _validate_docx_archive(data: bytes) -> None:
    if not data.startswith(b"PK\x03\x04"):
        raise DocumentProcessingError("The file does not contain a valid DOCX archive")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_FILES:
                raise DocumentProcessingError("The DOCX archive contains too many files")
            if sum(member.file_size for member in members) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentProcessingError("The DOCX expands beyond the safe processing limit")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentProcessingError("The DOCX structure is incomplete")
            for name in names:
                if name.startswith("/") or ".." in PurePosixPath(name).parts:
                    raise DocumentProcessingError("The DOCX contains an unsafe archive path")
    except BadZipFile as exc:
        raise DocumentProcessingError("The DOCX archive is invalid") from exc


def _extract_docx(data: bytes) -> str:
    _validate_docx_archive(data)
    try:
        document = Document(BytesIO(data))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    blocks.append(text)
        return "\n\n".join(blocks)
    except Exception as exc:
        raise DocumentProcessingError("The DOCX could not be read safely") from exc


class DocumentProcessor:
    """Base class for a pluggable format handler.

    Subclass, set `mime_by_extension` (extension -> MIME type) and implement
    `extract(data) -> str`, then call `register_processor(instance)`.
    """

    name: str = "base"
    mime_by_extension: dict[str, str] = {}

    def extract(self, data: bytes) -> str:  # pragma: no cover - interface
        raise NotImplementedError


class TextProcessor(DocumentProcessor):
    name = "text"
    mime_by_extension = {".txt": "text/plain", ".md": "text/markdown", ".markdown": "text/markdown"}

    def extract(self, data: bytes) -> str:
        return _extract_text(data)


class PdfProcessor(DocumentProcessor):
    name = "pdf"
    mime_by_extension = {".pdf": "application/pdf"}

    def extract(self, data: bytes) -> str:
        return _extract_pdf(data)


class DocxProcessor(DocumentProcessor):
    name = "docx"
    mime_by_extension = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }

    def extract(self, data: bytes) -> str:
        return _extract_docx(data)


_PROCESSORS: list[DocumentProcessor] = []
_PROCESSOR_BY_EXTENSION: dict[str, DocumentProcessor] = {}


def register_processor(processor: DocumentProcessor) -> None:
    """Add a format handler. Later formats (RTF, OCR/image, etc.) plug in here
    without changing extract_document(), the router, or the embedding flow."""
    _PROCESSORS.append(processor)
    for extension, mime in processor.mime_by_extension.items():
        _PROCESSOR_BY_EXTENSION[extension] = processor
        MIME_BY_EXTENSION[extension] = mime


def supported_extensions() -> list[str]:
    return list(MIME_BY_EXTENSION.keys())


# Default text-layer processors. OCR/image processors are intentionally absent.
register_processor(TextProcessor())
register_processor(PdfProcessor())
register_processor(DocxProcessor())


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    safe_filename = Path(filename or "").name
    extension = Path(safe_filename).suffix.lower()
    processor = _PROCESSOR_BY_EXTENSION.get(extension)
    if processor is None:
        labels = ", ".join(ext.lstrip(".").upper() for ext in supported_extensions())
        raise DocumentProcessingError(
            f"Supported formats are {labels}. Images/scanned documents need a future OCR processor."
        )
    if not data:
        raise DocumentProcessingError("The uploaded file is empty")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentProcessingError("The uploaded file exceeds the 5 MB limit")

    text = processor.extract(data)

    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentProcessingError("No extractable text was found; scanned documents require a future OCR processor")
    return ExtractedDocument(
        filename=safe_filename,
        file_type=MIME_BY_EXTENSION[extension],
        text=text,
        word_count=count_words(text),
    )